"""Check final EID DOCX artifacts structurally when visual rendering is unavailable."""

from __future__ import annotations

import argparse
from pathlib import Path
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "docs" / "submission_eid"

DOCX_FILES = {
    "manuscript": SUBMISSION / "manuscript_eid.docx",
    "cover_letter": SUBMISSION / "cover_letter_eid.docx",
    "author_statements": SUBMISSION / "author_statements.docx",
    "appendix": SUBMISSION / "supplements" / "Appendix_methods_eid.docx",
}

ASCII_MARKDOWN_FILES = [
    SUBMISSION / "manuscript_eid.md",
    SUBMISSION / "cover_letter_eid.md",
    SUBMISSION / "author_statements.md",
    SUBMISSION / "supplements" / "Appendix_methods_eid.md",
]


def _docx_xml(path: Path) -> str:
    with ZipFile(path) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def _has_forbidden_docx_parts(path: Path) -> bool:
    with ZipFile(path) as archive:
        names = set(archive.namelist())
    return any(
        name in names
        for name in (
            "word/comments.xml",
            "word/commentsExtended.xml",
            "word/people.xml",
            "word/revisions.xml",
        )
    )


def _margin_inches(value: int | None) -> float:
    return 0.0 if value is None else value / 1440.0


def _check_docx(name: str, path: Path) -> list[str]:
    errors: list[str] = []
    if not path.exists():
        return [f"Missing DOCX: {path}"]

    doc = Document(path)
    xml = _docx_xml(path)
    if _has_forbidden_docx_parts(path):
        errors.append(f"{name}: comments/revision parts present")
    if "<w:ins" in xml or "<w:del" in xml:
        errors.append(f"{name}: tracked change markup present")
    if "w:lnNumType" not in xml:
        errors.append(f"{name}: line numbering XML missing")

    for section in doc.sections:
        margins = [
            _margin_inches(section.top_margin.twips),
            _margin_inches(section.bottom_margin.twips),
            _margin_inches(section.left_margin.twips),
            _margin_inches(section.right_margin.twips),
        ]
        if any(abs(margin - 1.0) > 0.02 for margin in margins):
            errors.append(f"{name}: margins are not 1 inch: {margins}")

    if name == "manuscript":
        if len(doc.tables) < 3:
            errors.append("manuscript: expected at least 3 real Word tables")
        if len(doc.paragraphs) > 500:
            errors.append(
                "manuscript: paragraph count suggests wrapped Markdown lines were not merged"
            )
        if "Sparse Public Surveillance Limits Forecasting" not in "\n".join(
            paragraph.text for paragraph in doc.paragraphs[:20]
        ):
            errors.append("manuscript: title not found near document start")

    return errors


def _check_ascii_markdown(path: Path) -> list[str]:
    if not path.exists():
        return [f"Missing Markdown source: {path}"]
    text = path.read_text(encoding="utf-8")
    bad = sorted({char for char in text if ord(char) > 127})
    if bad:
        return [f"{path}: non-ASCII characters present: {bad}"]
    return []


def check_docx_artifacts() -> list[str]:
    errors: list[str] = []
    for name, path in DOCX_FILES.items():
        errors.extend(_check_docx(name, path))
    for path in ASCII_MARKDOWN_FILES:
        errors.extend(_check_ascii_markdown(path))
    return errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.parse_args()

    errors = check_docx_artifacts()
    if errors:
        print("EID DOCX artifact check: FAIL")
        for error in errors:
            print(f"- {error}")
        return 1
    print("EID DOCX artifact check: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
